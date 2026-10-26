import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r'E:/PWD301/PPTQ_LMS_FINAL_REPORT.docx'
md_path = r'E:/PWD301/docs/PPTQ_LMS_FINAL_REPORT.md'

print("Checking docx:", os.path.exists(doc_path))
doc = docx.Document(doc_path)
print("Paragraphs in docx:", len(doc.paragraphs))
print("Tables in docx:", len(doc.tables))

full_text = '\n'.join([p.text for p in doc.paragraphs])
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full_text += '\n' + cell.text

banned_words = ['Cisco', 'NetAcad', 'Coursera', 'Bloom', 'cognitive taxonomy', 'Xin lỗi vì đã làm bạn tức giận', 'hỏi ngoài lề']
found = {}
for bw in banned_words:
    count = full_text.lower().count(bw.lower())
    if count > 0:
        found[bw] = count

print("DOCX Banned words check (should be empty):", found)

keywords = ['PPTQ LMS', 'Bạch tuộc', 'Exam & Question Authoring Studio', 'Hỏi Bạch tuộc câu này', 'Lease', 'Fail-Closed', 'PBKDF2']
for kw in keywords:
    print(f'DOCX Keyword "{kw}": {full_text.lower().count(kw.lower())}')

with open(md_path, 'r', encoding='utf-8') as f:
    md_text = f.read()

found_md = {}
for bw in banned_words:
    count = md_text.lower().count(bw.lower())
    if count > 0:
        found_md[bw] = count

print("MD Banned words check (should be empty):", found_md)
for kw in keywords:
    print(f'MD Keyword "{kw}": {md_text.lower().count(kw.lower())}')
