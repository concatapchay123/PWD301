import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document(r'E:/PWD301/PPTQ_LMS_FINAL_REPORT.docx')
print("Total tables in doc:", len(doc.tables))

# Let's find the table that has 'pwd301_clamav'
for i, t in enumerate(doc.tables):
    for r in t.rows:
        row_text = [c.text.strip().replace('\n', ' ') for c in r.cells]
        if any('pwd301_clamav' in c for c in row_text):
            print(f"Found Table {i}: {len(t.rows)} rows, {len(t.columns)} cols")
            for row in t.rows:
                print("  Row:", [c.text.strip().replace('\n', ' ') for c in row.cells])
            break
