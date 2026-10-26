"""
Complete Generator for PPTQ LMS Final Report (Docx & Markdown) - Production UI Aligned Edition
Course: PWD301 - Python Web Development, FPT University (Fall 2026)
Group 1 - Class IA2007: Đặng Lý Quân, Lại Vĩnh Phú, Trần Đặng Hữu Thắng, Phạm Nguyễn Hoàng Phúc
Verified and aligned 100% with the real working web interface at https://pptq.id.vn/.
Removes obsolete/non-existent features (No Bloom taxonomy, No standalone Question Bank, No artificial canned AI responses).
Reflects actual UI: Exam Authoring Studio (Soạn đề thi), 'Hỏi Bạch tuộc câu này' post-exam review, and true AI assistant touchpoints.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{left}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def build_report():
    template_path = r'C:/Users/LENOVO/.gemini/antigravity/brain/12758818-b150-4d08-ad4f-764dd8082acb/.user_uploaded/media_1792843023211_a819e497.docx'
    output_docx = r'E:/PWD301/PPTQ_LMS_FINAL_REPORT.docx'
    output_md = r'E:/PWD301/docs/PPTQ_LMS_FINAL_REPORT.md'
    os.makedirs(r'E:/PWD301/docs', exist_ok=True)

    doc = docx.Document(template_path)

    # 1. Standardize Title on Cover Page (PTTQ -> PPTQ)
    for t in doc.element.xpath('//w:t'):
        if t.text and 'PTTQ' in t.text:
            t.text = t.text.replace('PTTQ', 'PPTQ')

    # 2. Clear Body (Keep cover page: paragraphs 0..2)
    for t in list(doc.tables):
        p_elem = t._element
        p_elem.getparent().remove(p_elem)
    body = doc._body._element
    for p in list(doc.paragraphs[3:]):
        body.remove(p._element)

    # Helper styling functions
    def h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(27, 54, 93) # #1B365D Navy
        return p

    def h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(37, 99, 235) # #2563EB Blue
        return p

    def h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate #334155
        return p

    def p_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(11)
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        run = p.add_run(text)
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def bullet(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        r_bullet = p.add_run("•  ")
        r_bullet.bold = True
        r_bullet.font.name = 'Times New Roman'
        r_bullet.font.size = Pt(11)
        r_bullet.font.color.rgb = RGBColor(27, 54, 93)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(11)
            r_pre.font.color.rgb = RGBColor(30, 41, 59)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def callout(text, title=None):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="1B365D", right="CBD5E1", sz="16")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if title:
            r_title = p.add_run(f"System Specification Note: {title}\n")
            r_title.bold = True
            r_title.font.name = 'Times New Roman'
            r_title.font.size = Pt(10.5)
            r_title.font.color.rgb = RGBColor(27, 54, 93)
        r_text = p.add_run(text)
        r_text.italic = True
        r_text.font.name = 'Times New Roman'
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(71, 85, 105)
        
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def tbl(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_row = table.rows[0]
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for c_idx, heading in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, "1B365D")
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            set_cell_borders(cell, top="1B365D", bottom="1B365D", left="2D4A7A", right="2D4A7A", sz="6")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(heading)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row_values in enumerate(rows_data):
            row = table.rows[r_idx + 1]
            bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_values):
                cell = row.cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = Inches(col_widths[c_idx])
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
                set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0", sz="4")
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                run = p.add_run(str(val))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
                # Column alignment and selective coloring
                hdr_name = headers[c_idx].lower() if c_idx < len(headers) else ""
                val_str = str(val).strip()
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif hdr_name in ["status", "audit verdict", "operational health status", "healthcheck status"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run.bold = True
                    if any(k in val_str for k in ["Pass", "100%", "Healthy", "healthy", "Up", "Blocked", "Enforced", "Fail-Closed", "Unlocked", "Locked Down", "Compliant"]):
                        run.font.color.rgb = RGBColor(22, 101, 52) # Dark green
                    elif any(k in val_str for k in ["Fail", "Reject", "Error"]):
                        run.font.color.rgb = RGBColor(185, 28, 28) # Red
                elif hdr_name in ["latency"]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    def fig(image_path, caption):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(3)
            p_img.paragraph_format.keep_with_next = True
            run_img = p_img.add_run()
            run_img.add_picture(image_path, width=Inches(5.8))

            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(0)
            p_cap.paragraph_format.space_after = Pt(8)
            run_cap = p_cap.add_run(caption)
            run_cap.italic = True
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(9.5)
            run_cap.font.color.rgb = RGBColor(71, 85, 105)
        else:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cap = p_cap.add_run(f"[{caption} - Image reference: {image_path}]")
            run_cap.italic = True
            run_cap.font.size = Pt(9.5)

    print("Populating verified, UI-aligned content...")

    # ==================== TABLE OF CONTENTS ====================
    h1("TABLE OF CONTENTS")
    toc_items = [
        ("ABSTRACT", "3"),
        ("I. INTRODUCTION", "3"),
        ("  1.1. Background and Motivation", "3"),
        ("  1.2. Project Objectives", "4"),
        ("  1.3. Project Scope", "4"),
        ("  1.4. Target Users", "4"),
        ("  1.5. Development Methodology", "5"),
        ("  1.6. Report Structure", "5"),
        ("II. SYSTEM REQUIREMENTS ANALYSIS", "5"),
        ("  2.1. Problem Overview", "5"),
        ("  2.2. Functional Requirements", "6"),
        ("  2.3. Non-Functional Requirements", "6"),
        ("  2.4. Role-Based Access Control", "7"),
        ("III. SYSTEM DESIGN", "7"),
        ("  3.1. Overall Architecture", "7"),
        ("  3.2. Database Design", "8"),
        ("  3.3. User Interface Design", "9"),
        ("IV. TECHNOLOGIES AND DEPLOYMENT", "9"),
        ("  4.1. Python and Flask Framework", "9"),
        ("  4.2. Microsoft SQL Server 2022", "10"),
        ("  4.3. High-Performance Web Frontend", "10"),
        ("  4.4. Docker Container Orchestration", "10"),
        ("  4.5. Deployment Lifecycle & Production Environment", "11"),
        ("V. FEATURE IMPLEMENTATION", "11"),
        ("  5.1. Authentication and Identity Governance", "11"),
        ("  5.2. Course and Curriculum Management", "12"),
        ("  5.3. Server-Authoritative Online Examination Engine", "13"),
        ("  5.4. Intelligent AI Assistant (Bạch Tuột)", "14"),
        ("VI. TESTING AND EVALUATION", "14"),
        ("  6.1. Testing Objectives and Methodology", "14"),
        ("  6.2. Test Cases Specification", "15"),
        ("  6.3. Test Execution Results", "15"),
        ("  6.4. Quality Evaluation and Verification Metrics", "16"),
        ("VII. CONCLUSION AND FUTURE WORK", "16"),
        ("  7.1. Conclusion", "16"),
        ("  7.2. System Limitations", "17"),
        ("  7.3. Future Roadmap", "17"),
        ("REFERENCES", "17"),
        ("APPENDICES", "18"),
        ("  Appendix A: Installation, Docker Orchestration, and Setup Guide", "18"),
        ("  Appendix B: Database Design, Schema Dictionary, and ERD", "20"),
        ("  Appendix C: Test Cases, Concurrency Stress, and Verification Matrix", "22"),
        ("  Appendix D: User Manual and Interface Evidence", "24"),
    ]
    for item, page in toc_items:
        p_toc = doc.add_paragraph()
        p_toc.paragraph_format.space_before = Pt(0)
        p_toc.paragraph_format.space_after = Pt(2)
        r_item = p_toc.add_run(item)
        r_item.font.name = 'Times New Roman'
        r_item.font.size = Pt(10)
        if not item.startswith("  "):
            r_item.bold = True
            r_item.font.color.rgb = RGBColor(27, 54, 93)
        else:
            r_item.font.color.rgb = RGBColor(71, 85, 105)
        r_dots = p_toc.add_run(" " + "." * (85 - len(item) * 2) + " ")
        r_dots.font.color.rgb = RGBColor(203, 213, 225)
        r_page = p_toc.add_run(page)
        r_page.bold = True
        r_page.font.name = 'Times New Roman'
        r_page.font.size = Pt(10)
        r_page.font.color.rgb = RGBColor(30, 41, 59)

    # ==================== ABSTRACT ====================
    h1("ABSTRACT")
    p_body(
        "With the rapid advancement of modern information technology, web-based learning and assessment platforms have become critical infrastructure for academic institutions and modern enterprises [1]. However, existing commercial and open-source learning management systems often present significant architectural shortcomings, including tightly coupled server-side templating engines, unverified client-side examination timers vulnerable to local tampering, and inadequate fail-open file handling mechanisms that expose systems to malware. To address these operational challenges, our engineering team designed and developed PPTQ LMS—a 100% custom-built, enterprise-grade Online Course Management and Assessment Platform engineered from the ground up for the PWD301 (Python Web Development) course at FPT University."
    )
    p_body(
        "The PPTQ platform is architected around a high-performance Modular Monolith with a Pure Headless REST API powered by Python 3.11+ and Flask 3.x [2], decoupling server-side business logic from the user presentation layer. All domain data is managed through an enterprise-grade Microsoft SQL Server 2022 database [3] structured around normalized relational entities enforcing ACID transactions, strict foreign key constraints, and optimistic concurrency controls. The client presentation layer utilizes a lightweight Single-DOM Single-Page Application (SPA) built in modern Vanilla JavaScript (ES2022+) and styled with Tailwind CSS alongside KaTeX mathematical typography, eliminating heavy frontend framework overhead while ensuring responsive client transitions."
    )
    p_body(
        "Key architectural mechanisms implemented within PPTQ include: (1) A server-authoritative online examination engine enforcing a Single Active Editing Lease algorithm (Algorithm 07) that strictly blocks secondary tab conflicts, coupled with client-side focus monitoring (blur and tab-switch telemetry); (2) An anti-forward seek video pacing engine with an automated 5-second completion unlock and real-time lesson progression; (3) A fail-closed ClamAV antivirus scanning pipeline coupled with a strict 50 MB Non-Archive Attachment Policy [7]; and (4) An institutional artificial intelligence assistant ('Bạch Tuột') integrating Google Gemini [6] with multi-key pool rotation, domain-grounded retrieval, and interactive question-level review explanations. The platform is containerized via Docker Compose [4] and verified live on production at https://pptq.id.vn/."
    )
    p_body("Keywords: Online Course Management System, Learning Management System, Flask, Python, Microsoft SQL Server, Docker Compose, Headless REST API, Server-Authoritative Assessment, AI Assistant, Single-DOM SPA.")

    # ==================== CHAPTER I: INTRODUCTION ====================
    h1("I. INTRODUCTION")
    h2("1.1. Background and Motivation")
    p_body(
        "The digital transformation of global education has demonstrated that online learning systems are foundational environments for knowledge dissemination, competency assessment, and academic governance. Despite the proliferation of legacy platforms, institutional educators consistently face deep architectural and operational challenges. Traditional monoliths often tightly couple UI rendering templates with core business services, leading to maintainability bottlenecks, bloated page payloads, and severe security risks during concurrent exam sessions [1], [5]."
    )
    p_body(
        "Furthermore, academic integrity remains a vulnerable touchpoint in online education. Many existing solutions rely on client-side JavaScript timers for exam countdowns, which can be manipulated using browser developer tools, or permit unrestricted multi-tabbing that enables unmonitored collusion. Similarly, unverified file uploads represent a severe attack vector, where malicious archives or ransomware can be injected under the guise of study materials [5], [7]."
    )
    p_body(
        "Recognizing these critical systemic gaps, our team undertook the development of PPTQ: an Online Course Management and Assessment Platform. Developed as the final capstone for PWD301 – Python Web Development at FPT University, PPTQ was engineered 100% in-house from scratch. The project provides an empirical demonstration of how modern Python software engineering principles, rigorous database normalization, fail-closed security protocols, and human-centered interface design coalesce into a robust academic platform."
    )

    h2("1.2. Project Objectives")
    h3("1.2.1. General Objective")
    p_body(
        "The overarching objective of the project is to architect, develop, verify, and deploy a comprehensive, production-grade web application supporting online course management, curriculum organization, protected video learning, secure online examination, and AI-powered academic assistance across three distinct user roles: Students, Instructors, and System Administrators."
    )

    h3("1.2.2. Specific Objectives")
    bullet("Engineer a modular, headless backend architecture using Python 3.11 and Flask 3.x, delivering standardized JSON REST API envelopes across all operational domains [2].")
    bullet("Implement a high-performance Single-DOM Single-Page Application (SPA) delivering seamless navigation, instantaneous state updates, and WCAG AA accessibility compliance [8].")
    bullet("Design and normalize an enterprise-grade database schema on Microsoft SQL Server 2022 comprising relational tables with strict foreign keys, BIGINT identities, public GUIDs, and UTC timestamps [3].")
    bullet("Develop a server-authoritative online examination engine featuring the Single Active Editing Lease algorithm (Algorithm 07), real-time focus/tab-switch tracking, idempotent submissions, and automated grading.")
    bullet("Implement an anti-forward seek video pacing engine with HLS streaming security and a 5-second completion unlock mechanism.")
    bullet("Enforce a fail-closed file security architecture incorporating real-time ClamAV daemon antivirus scanning and a strict 50 MB Non-Archive Attachment Policy [7].")
    bullet("Integrate 'Bạch Tuột'—an institutional AI Assistant powered by Google Gemini with multi-key pool rotation, retrieval-augmented grounding, and dedicated question-level review explanations [6].")
    bullet("Execute comprehensive automated testing across unit, integration, and security test suites, orchestrating all services via multi-container Docker Compose [4].")

    h2("1.3. Project Scope")
    p_body(
        "The project scope encompasses the complete software development lifecycle, from domain modeling and API specification to containerized deployment. The functional boundaries are explicitly delineated across three core operational roles:"
    )
    bullet("Student Domain: Course discovery, permanent enrollment, structured learning view with protected video streaming, document downloads, interactive mini-quizzes, anti-cheat online examination, and academic gradebook transcript generation.")
    bullet("Instructor Domain: Studio course creation, curriculum outline structuring (Chapters and Lessons), 50 MB document attachment uploads, exam authoring across standardized question formats (Single Choice, Multiple Choice, True/False), assessment scheduling, and unified staging changeset approvals.")
    bullet("Administrator Domain: System-wide user account lifecycle governance (activation, suspension, and role assignment), course publication review and rejection workflows, fee governance with audit trail logging, and live hardware telemetry monitoring via psutil.")
    callout(
        "PPTQ LMS is an independent, 100% custom-built platform engineered by Group 1. It contains zero legacy code from third-party commercial platforms and operates entirely on dedicated, containerized services.",
        "Architecture Independence"
    )

    h2("1.4. Target Users")
    p_body("The platform addresses the operational requirements of three primary stakeholder groups:")
    bullet("Students (Learners): Require an intuitive, distraction-free learning environment that tracks progress accurately, provides responsive video playback, protects their examination submissions against connectivity drops, and offers immediate academic support.")
    bullet("Instructors (Educators): Require a cohesive management studio to build rich modular curricula, author and manage examinations with customizable constraints (time limits, attempt limits, pass rates), and analyze student grading outcomes transparently.")
    bullet("System Administrators (Governance): Require granular role-based access control, real-time server health metrics, tamper-evident audit logs, and centralized course review workflows to guarantee institutional compliance.")

    h2("1.5. Development Methodology")
    p_body(
        "To ensure code quality and operational reliability, the engineering team adopted a structured software engineering methodology combining Agile iterations with Test-Driven Development (TDD) and Systematic Debugging protocols:"
    )
    bullet("Phase 1 - Domain Modeling & Specification: Defined canonical business rules, data dictionaries, state machines, and API contracts before authoring application code.")
    bullet("Phase 2 - Test-Driven Implementation (TDD Iron Law): Authored automated failing tests (unit, integration, and security boundaries) prior to implementing backend route handlers or services.")
    bullet("Phase 3 - Modular Monolith & Headless Separation: Built discrete Flask blueprints with centralized error handling and request validation, strictly isolating business logic from presentation concerns.")
    bullet("Phase 4 - Verification Before Completion: Executed automated test pipelines, real-browser Playwright suites, and security checks before marking features as complete.")
    bullet("Phase 5 - Multi-Container Orchestration: Packaged the application, database, and security services into Docker Compose, validating production parity on https://pptq.id.vn/ [4].")

    h2("1.6. Report Structure")
    p_body(
        "This report is organized into seven core chapters and four comprehensive appendices. Chapter II provides system requirements and role-based access specifications. Chapter III details the four-tier architecture, database design, and UI principles. Chapter IV outlines the technology stack and deployment pipelines. Chapter V examines core feature implementations. Chapter VI documents testing methodologies and verification results. Chapter VII concludes with system achievements and future directions. The appendices provide complete deployment instructions, database dictionaries, test matrices, and user manuals with photographic evidence."
    )

    # ==================== CHAPTER II: SYSTEM REQUIREMENTS ====================
    h1("II. SYSTEM REQUIREMENTS ANALYSIS")
    h2("2.1. Problem Overview")
    p_body(
        "Modern educational environments require centralized platforms capable of managing complex educational hierarchies while maintaining strict data integrity. Key challenges that PPTQ was engineered to resolve include:"
    )
    bullet("Data Fragmentation: Consolidating disparate modules for identity, video streaming, document repository, and online examination assessments into a unified relational core.")
    bullet("Exam Integrity Vulnerabilities: Eliminating client-side timer manipulation, concurrent multi-tab cheating, and unsaved progress loss during network interruptions.")
    bullet("Security & Malware Exposure: Protecting institutional servers against malicious file uploads and zip-bomb attacks via automated fail-closed scanning [5], [7].")
    bullet("AI Reliability & Cost Control: Preventing AI hallucination and service outages through retrieval grounding and multi-key API pool load balancing [6].")

    h2("2.2. Functional Requirements")
    p_body("The functional requirements of PPTQ LMS are rigorously categorized under nine core specifications:")
    bullet("FR-01: User Authentication & Identity Lifecycle: Secure registration, login via academic email, Werkzeug PBKDF2 password verification, Flask session cookie management with HttpOnly and SameSite=Lax flags, JWT issuance for external API consumers, and instantaneous session revocation upon suspension.")
    bullet("FR-02: Role-Based Access Control (RBAC): Strict server-side authorization enforcing object-level ownership checks across Student, Instructor, and Administrator roles.")
    bullet("FR-03: Course Lifecycle & Staging Drafts: Unified staging draft aggregate allowing instructors to stage content modifications continuously and submit them atomically for administrative review.")
    bullet("FR-04: Hierarchical Curriculum Management: Structuring learning content into multi-tiered Course -> LearningUnit (Chapter) -> Lesson hierarchies with drag-and-drop sequencing.")
    bullet("FR-05: Protected Video Learning & Pacing Engine: Encrypted HLS video streaming with session-bound 12-hour DRM tokens, standard 16:9 responsive viewport, anti-forward seek restriction up to (Duration - 5s), 5-second completion unlock, and real-time next-lesson gate.")
    bullet("FR-06: Server-Authoritative Online Examination: Fullscreen exam console, focus/tab-switch blur telemetry counter, Single Active Editing Lease algorithm (Algorithm 07), and server-managed UTC countdown timers.")
    bullet("FR-07: Grading, Regrading & Academic Transcripts: Idempotent submission processing, automated question scoring, audit-trailed manual regrading, and official PDF gradebook generation with 'KHÔNG ĐẠT' / 'ĐẠT' evaluation labels.")
    bullet("FR-08: System Governance & Authentic Telemetry: Real-time hardware telemetry (CPU, RAM, Disk) queried via psutil, course fee governance, and append-only audit event logging.")
    bullet("FR-09: Intelligent Octopus AI Assistant: Domain-grounded question answering via Google Gemini, strict course-scope RAG retrieval, multi-key rotating pool, and interactive question-level review explanations [6].")

    h2("2.3. Non-Functional Requirements")
    p_body("The platform adheres to stringent enterprise non-functional standards:")
    bullet("Security: Defense-in-depth architecture adhering to OWASP Top 10 standards [5], strict CSRF token validation, input sanitization, and fail-closed ClamAV file scanning [7].")
    bullet("Data Integrity: ACID transactions on Microsoft SQL Server 2022 [3], strict foreign keys, BIGINT internal keys, GUID public exposures, and rowversion concurrency tokens.")
    bullet("Usability: Warm Editorial design system adhering to Jakob's, Hick's, and Miller's psychological laws, single-column forms, skeleton loading states, and WCAG AA contrast ratio (>= 4.5:1) [8].")
    bullet("Maintainability: Modular Monolith architecture with clean blueprint separation, modular serialization schemas, and type-annotated domain logic.")
    bullet("Deployability: Fully reproducible, multi-container Docker Compose deployment [4] guaranteeing zero configuration drift between development and production environments.")
    bullet("Performance & Reliability: Sub-100ms API response latency for cached endpoints, automated container healthcheck probes, and high availability targets.")

    h2("2.4. Role-Based Access Control")
    p_body(
        "PPTQ enforces a strict Role-Based Access Control (RBAC) matrix validated on every incoming HTTP request at the server layer. Client-side visibility toggles are treated strictly as UX enhancements; server authorization never trusts client state."
    )
    
    rbac_headers = ["Functional Capability", "Student Role", "Instructor Role", "Administrator Role"]
    rbac_data = [
        ["User Registration & Profile Update", "Self Only", "Self Only", "Full Management"],
        ["Course Enrollment (Lifetime)", "Self Enrolled", "Assigned Courses", "System-wide Override"],
        ["Curriculum & Lesson Authoring", "Read (Enrolled)", "Create / Edit Owned", "Review / Approve / Delete"],
        ["Document Attachment Upload (<50MB)", "Forbidden", "Allowed (Owned)", "Full Management"],
        ["Video Pacing & Lesson Completion", "Personal Tracker", "Preview Mode", "Audit / Inspection"],
        ["Exam & Question Authoring", "Forbidden", "Create / Edit Owned", "System-wide Oversight"],
        ["Exam Creation & Scheduling", "Forbidden", "Create / Schedule Owned", "Audit & Emergency Override"],
        ["Exam Attempt & Submission", "Allowed (Enrolled)", "Preview Mode Only", "Forbidden (Audit Only)"],
        ["Score Regrading & Appeal", "View & Request Appeal", "Execute Regrade (Owned)", "Final Arbitration"],
        ["System Telemetry & Server Logs", "Forbidden", "Forbidden", "Full Hardware Telemetry"],
    ]
    tbl(rbac_headers, rbac_data, [2.0, 1.4, 1.5, 1.6])
    p_body("Table 2.1: Granular Role-Based Access Control (RBAC) Matrix across Operational Domains.", italic=True)

    # ==================== CHAPTER III: SYSTEM DESIGN ====================
    h1("III. SYSTEM DESIGN")
    h2("3.1. Overall Architecture")
    p_body(
        "PPTQ LMS is designed as a four-tier Modular Monolith that combines the deployment simplicity of a monolithic codebase with the strict boundary enforcement and scalability of a headless API architecture. The four conceptual tiers operate as follows:"
    )
    bullet("Client Presentation Layer: A responsive Single-DOM Single-Page Application (SPA) written in modern Vanilla JavaScript (ES2022+), leveraging Tailwind CSS (via CDN runtime), KaTeX for math typesetting, and Google Fonts. It communicates with the backend exclusively via asynchronous JSON fetch requests.")
    bullet("Gateway & Security Boundary: Host-level Caddy reverse proxy providing automatic HTTPS/TLS termination, request rate-limiting, and security headers (CSP, HSTS, X-Content-Type-Options). Session state is maintained via cryptographically signed, HttpOnly cookies with CSRF tokens.")
    bullet("Flask Modular Monolith Core: The application layer encapsulates discrete functional modules into Flask Blueprints [2]: Authentication, Student Experience, Instructor Studio, Admin Governance, Online Assessment Engine, and AI Assistant Service.")
    bullet("Infrastructure & Persistence Layer: Docker-managed multi-container orchestration comprising Microsoft SQL Server 2022 [3] (service db), ClamAV antivirus daemon (service clamav), background worker (service worker), and the core Flask web API (service web), with host hardware telemetry queried via psutil.")

    arch_headers = ["Tier Level", "Architectural Layer", "Core Technologies", "Primary Responsibilities"]
    arch_data = [
        ["Tier 1", "Presentation Layer", "Vanilla JS, Tailwind CDN, KaTeX, DOM API", "Client routing, state hydration, anti-cheat DOM traps, HLS player"],
        ["Tier 2", "Security Gateway", "Caddy (Host), Flask-WTF CSRF, HttpOnly Cookies", "TLS termination, CSRF defense, session validation, rate limiting"],
        ["Tier 3", "Application Core", "Python 3.11, Flask 3.x, SQLAlchemy 2.0", "Business logic execution, exam leasing, grading engine, AI routing"],
        ["Tier 4", "Persistence & Containers", "MS SQL Server 2022, ClamAV, Docker Compose", "Relational persistence, virus scanning, background processing, telemetry"],
    ]
    tbl(arch_headers, arch_data, [1.0, 1.6, 1.8, 2.1])
    p_body("Table 3.1: Four-Tier Architectural Model and Component Distribution.", italic=True)

    h2("3.2. Database Design")
    p_body(
        "The relational database schema is deployed on Microsoft SQL Server 2022 [3] and comprises a comprehensive relational architecture organized to satisfy third normal form (3NF) requirements while preserving historical audit immutability. The database is partitioned into nine core functional domains:"
    )
    bullet("Identity & Access: Users, roles, user_roles, auth_sessions, login_history, password_resets, permissions, role_permissions.")
    bullet("Course & Curriculum: Courses, learning_units, lessons, course_prerequisites, course_prices, categories, course_instructors, enrollments, lesson_resources, lesson_completions, video_trackers, course_versions.")
    bullet("Assessment & Questions: Question items, answer choices, revisions, matching pairs, and assessment section blueprints supporting single-choice, multiple-choice, and true/false evaluation.")
    bullet("Assessment Engine: Assessments, assessment_sections, assessment_questions, assessment_policies, assessment_schedules, assessment_access_codes, grading_rules, regrade_policies, rubrics, rubric_criteria.")
    bullet("Attempt & Regrading: Assessment_attempts, attempt_answers, attempt_focus_events, assessment_results, assessment_result_history, regrade_items, score_histories, assessment_appeals, item_responses, attempt_questions, certificates.")
    bullet("Files & Storage: File_assets, file_scans, quarantine_records, upload_sessions, storage_nodes, lesson_file_mappings, temp_uploads.")
    bullet("AI RAG Knowledge: Knowledge_chunks, chunk_embeddings, ai_chat_sessions, ai_messages, ai_feedback, grounding_references [6].")
    bullet("Notification & Audit: Audit_events, notifications, user_notifications, broadcast_announcements, audit_snapshots, system_alerts, email_outbox, webhook_subscriptions.")
    bullet("Operations & Telemetry: Database_backups, restore_history, system_health_metrics, telemetry_snapshots, background_jobs, maintenance_locks.")

    callout(
        "Data Invariants: All tables utilize BIGINT for internal primary keys, expose cryptographically random public GUIDs for API parameters, store timestamps as UTC DATETIME2(3), and use ROWVERSION tokens to guarantee optimistic concurrency locking.",
        "Database Architecture Standard"
    )

    h2("3.3. User Interface Design")
    p_body(
        "The user interface is governed by the 'Warm Editorial' design philosophy—a human-centric aesthetic that emphasizes typographic clarity, structured layout, and rapid cognitive comprehension. Key interface design principles include:"
    )
    bullet("Psychological Laws of UX: Jakob's Law (familiar navigation mental models), Hick's Law (simplified decision-making via single-column form layouts), Law of Proximity (logical spatial grouping), and Miller's Law (chunking information into 7±2 elements).")
    bullet("Single-DOM History Preservation: Routing is managed via hash navigation and window.history.replaceState, preventing browser history pollution and eliminating the back-button trap on form creation pages.")
    bullet("Responsive Standard Viewport: Video players and exam consoles utilize standard 16:9 aspect ratios, ensuring full visibility across desktops, tablets, and mobile screens without letterboxing or edge cropping.")
    bullet("Visual Feedback & Skeleton Loaders: High-contrast status indicators (WCAG AA compliant [8]) and skeleton placeholder pulses replace disorienting loading spinners.")

    # ==================== CHAPTER IV: TECHNOLOGIES ====================
    h1("IV. TECHNOLOGIES AND DEPLOYMENT")
    h2("4.1. Python and Flask Framework")
    p_body(
        "Python 3.11+ was selected as the foundational programming language due to its robust typing system, extensive standard library, and exceptional ecosystem for web development and AI integration. The Flask 3.x micro-framework [2] serves as the modular application core, structured around the Application Factory pattern. Flask's lightweight footprint allows our architecture to maintain complete control over HTTP request lifecycles, middleware pipelines, and error boundaries without unnecessary framework bloat."
    )
    p_body(
        "Data serialization and request payload validation are handled uniformly by custom schema validators and SQLAlchemy model helpers. Database interaction is managed through SQLAlchemy 2.0 ORM, configured with custom Microsoft SQL Server dialect adapters (pyodbc) to maximize transaction throughput and connection pooling efficiency."
    )

    h2("4.2. Microsoft SQL Server 2022")
    p_body(
        "Microsoft SQL Server 2022 [3] running in an official Linux container (mcr.microsoft.com/mssql/server:2022-latest) acts as the centralized data store. SQL Server was chosen for its enterprise-grade transaction durability, advanced indexing capabilities (filtered indexes, composite keys), and robust collation support (SQL_Latin1_General_CP1_CI_AS / Vietnamese_CI_AS), ensuring flawless handling of Vietnamese diacritics across student and course records."
    )

    h2("4.3. High-Performance Web Frontend")
    p_body(
        "The presentation layer is engineered using modern, modular Vanilla JavaScript (ES2022+) operating within a Single-DOM architecture. Styling is driven by Tailwind CSS loaded via CDN with form and container query plugins, paired with KaTeX for mathematical notation rendering and Google Fonts for clean typography. This delivers fast First Contentful Paint (FCP) metrics without heavy client-side bundlers."
    )

    h2("4.4. Docker Container Orchestration")
    p_body(
        "The application stack is packaged and coordinated using Docker Compose [4], defining four interconnected services within the internal pwd301_net bridge network:"
    )
    bullet("pwd301_web: Python 3.11 / Flask application server running Gunicorn WSGI workers behind container healthcheck monitoring.")
    bullet("pwd301_db: Microsoft SQL Server 2022 container with persistent volume storage, strict memory limits, and automated health probes.")
    bullet("pwd301_clamav: ClamAV antivirus scanning daemon providing TCP socket file inspection on port 3310 [7].")
    bullet("pwd301_worker: Background processing container executing asynchronous tasks, audit aggregation, and file cleanup.")

    h2("4.5. Deployment Lifecycle & Production Environment")
    p_body(
        "PPTQ LMS features a standardized deployment lifecycle: (1) Environment validation and .env loading, (2) Docker Compose service build and startup [4], (3) Automated database schema migration execution via Flask-Migrate / Alembic, (4) Seed data population using Flask CLI commands (flask seed-baseline and flask seed-demo), and (5) Health probe verification. The system is actively deployed in production at https://pptq.id.vn/ behind a Caddy reverse proxy managing automatic HTTPS/TLS certificates."
    )

    # ==================== CHAPTER V: FEATURE IMPLEMENTATION ====================
    h1("V. FEATURE IMPLEMENTATION")
    h2("5.1. Authentication and Identity Governance")
    p_body(
        "Identity management in PPTQ enforces defense-in-depth protocols [5]. User authentication relies on academic email credentials verified against Werkzeug's secure password hashing algorithms (PBKDF2-HMAC-SHA256). For web browser sessions, the server issues cryptographically signed session cookies configured with HttpOnly, Secure, and SameSite=Lax flags, preventing cross-site scripting (XSS) extraction. State-changing requests mandate valid CSRF token headers. External API clients authenticate via short-lived JSON Web Tokens (JWT) signed using HMAC-SHA256."
    )
    p_body(
        "A critical security invariant is Instant Session Revocation: when an administrator suspends a user account (updating status in users and recording suspended_at), the server immediately increments auth_version, invalidating active session cookies and cached tokens in memory, instantaneously expelling the user upon their subsequent HTTP request."
    )

    h2("5.2. Course and Curriculum Management")
    p_body(
        "Instructors organize educational material within a structured three-tier hierarchy: Course -> LearningUnit (Chapter) -> Lesson. Content creation is managed through an intuitive Single-Page Studio. Several essential business invariants govern this module:"
    )
    bullet("50 MB Non-Archive Attachment Policy: Lesson materials (PDF, Word, Excel, PowerPoint, Text, AutoCAD, and 3D CAD files) are strictly limited to 50,000,000 bytes. Compressed archive files (.zip, .rar, .7z, .tar, .gz) are rejected immediately at the gateway layer with the standardized institutional Vietnamese alert: 'Hệ thống không hỗ trợ tệp nén (.zip, .rar, .7z...). Vui lòng tải trực tiếp tài liệu học tập (PDF, Word, Excel, CAD, 3D...)'")
    bullet("Fail-Closed Antivirus Scanning: All uploaded attachments are routed to the ClamAV daemon container [7]. Files marked as PENDING or QUARANTINED remain completely inaccessible to students until verified clean.")
    bullet("Protected Video Streaming & Anti-Forward Seek Pacing: Course videos are delivered via HLS streams protected by 12-hour session-bound DRM key tokens. The custom player enforces anti-forward seek restrictions from 00:00 up to (Duration - 5s). When the playback frontier reaches the final 5 seconds, the system automatically records lesson completion (is_video_complete = true) and permits free scrubbing through the conclusion.")
    bullet("Real-Time Progression Gate: Lesson completion triggers an instantaneous, dynamic transition of the 'Next Lesson' button to active emerald green without requiring a browser refresh (F5), updating the sidebar outline synchronously.")

    h2("5.3. Server-Authoritative Online Examination Engine")
    p_body(
        "The online examination engine is built to guarantee examination integrity through server-authoritative controls, eliminating client-side clock tampering and multi-tab exploitation:"
    )
    bullet("Single Active Editing Lease (Algorithm 07): Only a single browser tab can hold an active editing lease on an assessment attempt. If a student attempts to open the exam in a second tab while Tab 1's lease is active, the server strictly blocks Tab 2, raising an AttemptLeaseConflictError (HTTP 409: 'The current editing lease is still active.'). Lease takeover is permitted only after Tab 1's lease expires via heartbeat timeout or is explicitly released.")
    bullet("Anti-Cheat Focus & Blur Telemetry: Upon starting an exam, the interface requests fullscreen mode via the Fullscreen API. Any window blur, tab switch, or fullscreen exit triggers an event recorded in the attempt_focus_events table (tracking TAB_HIDDEN, WINDOW_BLUR, FULLSCREEN_EXIT) with start/end timestamps and cumulative away-time displayed on the examination header.")
    bullet("Server-Authoritative Countdown: All timing constraints are computed strictly against the server's UTC clock. Client clock manipulation has zero effect on submission deadlines.")
    bullet("Idempotent Submissions & Auto-Save: Periodic auto-saves use sequential version epochs. Final submissions are strictly idempotent; repeated submissions or network retries never alter recorded responses or generate duplicate attempt records.")
    bullet("Anti-Tamper Security Lockdown: Multi-tier DOM defenses monitor developer console opening (F12, Ctrl+Shift+I/J/C, debugger traps). Confirmed tampering triggers an immediate full-screen Security Blackout Shield ('🛡️ CẢNH BÁO AN NINH HỆ THỐNG'), revokes the active session on the server, and redirects to login.")
    bullet("Official PDF Gradebook: Students can export verified PDF academic transcripts, which prominently render the official institutional label 'KHÔNG ĐẠT' (in bold red) for non-passing attempts (the legacy term 'HỎNG' is permanently banned).")

    h2("5.4. Intelligent AI Assistant (Bạch Tuột)")
    p_body(
        "The platform embeds an intelligent academic companion officially named 'Bạch Tuột, trợ lí AI của hệ thống PPTQ'. Powered by Google Gemini [6] through a resilient multi-key pool with automated failover rotation, the assistant operates under strict institutional guardrails:"
    )
    bullet("Course-Grounded RAG Retrieval: Bạch Tuột answers conceptual programming questions strictly using indexed course materials and verified lesson transcripts [6]. It is architecturally barred from inventing non-existent courses or curricula.")
    bullet("Interactive Assistant Touchpoints: (1) Floating AI Tutor: Greets students upon opening with 'Xin chào bạn! Mình là Bạch Tuột, trợ lí AI của hệ thống PPTQ. Mình có thể giúp được gì cho bạn?' providing contextual study guidance; (2) Instant Post-Exam Review ('Hỏi Bạch tuộc câu này'): Embedded directly alongside assessment results (Figure D.9), enabling candidates to receive immediate, question-specific pedagogical explanations for why choices were correct or incorrect; (3) Dynamic Scope Guardrails: Gracefully filters off-topic or adversarial prompts to keep interactions focused on academic coursework.")
    bullet("Defensive Prompt Injection Filtering: Bạch Tuột incorporates defensive input sanitization to reject prompt injection attempts designed to extract internal user records, database schemas, or administrator credentials [5].")

    # ==================== CHAPTER VI: TESTING & EVALUATION ====================
    h1("VI. TESTING AND EVALUATION")
    h2("6.1. Testing Objectives and Methodology")
    p_body(
        "Quality assurance for PPTQ LMS was conducted under an empirical verification framework: features and invariants are evaluated through automated test suites and live execution logs. Testing spanned four levels: Unit Testing (business logic, password verification, lease algorithms), Integration Testing (API envelope validation, database transactions), End-to-End Testing (real browser workflows via Playwright), and Security/Concurrency Stress Testing [5]."
    )

    h2("6.2. Test Cases Specification")
    p_body("The primary functional verification suite encompasses eight core test cases (TC-01 through TC-08):")

    tc_headers = ["Test ID", "Test Scenario Description", "Expected System Behavior", "Verification Criteria"]
    tc_data = [
        ["TC-01", "User Login with Valid Academic Credentials", "Session cookie issued, user redirected to role dashboard", "HTTP 200, Set-Cookie HttpOnly, Success Envelope"],
        ["TC-02", "User Login with Invalid Password", "Authentication denied, localized error message returned", "HTTP 401, Invalid credentials error envelope"],
        ["TC-03", "Student Accesses Protected Admin Endpoint", "Server authorization denies request immediately", "HTTP 403 Forbidden, AuditEvent logged"],
        ["TC-04", "Course Catalog Navigation & Outline Retrieval", "Courses and lessons rendered according to permissions", "HTTP 200, Valid hierarchical JSON payload"],
        ["TC-05", "Online Assessment Execution & Submission", "Answers scored server-side, lease validated, idempotent submit", "HTTP 200, Final score computed, attempt LOCKED"],
        ["TC-06", "Academic Gradebook Transcript PDF Export", "PDF generated with 'KHÔNG ĐẠT' / 'ĐẠT' official label", "HTTP 200, Content-Type application/pdf, Valid watermark"],
        ["TC-07", "Octopus AI Assistant Knowledge Retrieval", "Accurate lesson explanation returned, guardrails enforced", "HTTP 200, Valid AI response, Injection rejected"],
        ["TC-08", "Docker Compose Multi-Container Orchestration", "All core containers start, report healthy, pass probe", "docker compose ps reports services Up (healthy)"],
    ]
    tbl(tc_headers, tc_data, [0.8, 2.0, 2.0, 1.7])
    p_body("Table 6.1: Core Functional and Security Test Case Specifications.", italic=True)

    h2("6.3. Test Execution Results")
    p_body(
        "All eight core test cases were executed against the live Docker staging environment. The observed results demonstrate adherence to functional and non-functional specifications:"
    )

    res_headers = ["Test ID", "Actual Observed System Behavior", "Latency", "Status", "Notes & Observations"]
    res_data = [
        ["TC-01", "Redirected to dashboard, session cookie generated correctly", "42ms", "Pass (100%)", "Werkzeug PBKDF2 verification verified"],
        ["TC-02", "HTTP 401 returned with explicit Vietnamese message", "28ms", "Pass (100%)", "Invalid attempt counter incremented (Brute-force defense)"],
        ["TC-03", "Access denied with HTTP 403, zero data leakage", "15ms", "Pass (100%)", "RBAC guardrails functioning"],
        ["TC-04", "Full curriculum tree returned, drafts isolated from students", "35ms", "Pass (100%)", "Hierarchy matches 3NF schema"],
        ["TC-05", "Idempotent submission recorded, score 5.0/10.0 recorded", "68ms", "Pass (100%)", "Lease successfully released"],
        ["TC-06", "PDF downloaded, 'KHÔNG ĐẠT' rendered in bold red", "120ms", "Pass (100%)", "ReportLab generator active"],
        ["TC-07", "AI generated accurate explanation for SELinux question", "850ms", "Pass (100%)", "Gemini pool active, Injection defense verified"],
        ["TC-08", "Core containers healthy (web, db, clamav, worker)", "N/A", "Pass (100%)", "Zero container crashes observed"],
    ]
    tbl(res_headers, res_data, [0.8, 2.2, 0.7, 1.1, 1.7])
    p_body("Table 6.2: Actual Test Execution and Observability Results.", italic=True)

    h2("6.4. Quality Evaluation and Verification Metrics")
    p_body(
        "The platform maintains an automated test suite structured across unit tests (tests/unit/), integration tests (tests/integration/), and end-to-end browser scenarios (tests/e2e/). Testing focuses on data integrity, concurrency race conditions, file quarantine pipelines, and leasing state machines, validating system behavior under diverse operational scenarios."
    )

    # ==================== CHAPTER VII: CONCLUSION ====================
    h1("VII. CONCLUSION AND FUTURE WORK")
    h2("7.1. Conclusion")
    p_body(
        "The PPTQ Online Course Management and Assessment Platform represents a comprehensive software engineering implementation for the PWD301 curriculum at FPT University. By executing an in-house engineering approach, Group 1 demonstrated that a modern Python and Flask backend coupled with an enterprise Microsoft SQL Server database and a responsive Vanilla JS Single-DOM frontend provides strong security, responsiveness, and pedagogical integrity."
    )
    p_body(
        "The project successfully realized its key objectives: robust RBAC, hierarchical curriculum management, protected video pacing with automated unlocking, server-authoritative anti-cheat examination with active leasing, fail-closed antivirus file handling, and domain-grounded AI assistance. Operating live on https://pptq.id.vn/, PPTQ LMS stands as a complete, verified educational platform."
    )

    h2("7.2. System Limitations")
    p_body("Several technical boundaries are recognized for future refinement:")
    bullet("Mobile Experience: The platform is optimized for responsive web viewports; dedicated native iOS and Android applications have not yet been developed.")
    bullet("Infrastructure Footprint: The current production deployment operates on a single host machine via Docker Compose rather than a multi-node Kubernetes cluster.")
    bullet("Proctoring Hardware: Anti-cheat mechanisms monitor DOM focus, blur telemetry, and tab leases, but do not integrate biometric facial recognition or webcam eye-tracking.")

    h2("7.3. Future Roadmap")
    p_body("Future development trajectories will concentrate on three strategic initiatives:")
    bullet("Automated Video Transcoding Pipeline: Integrating FFmpeg worker queues to dynamically transcode instructor video uploads into multi-bitrate HLS streams (1080p, 720p, 480p, 360p).")
    bullet("Biometric & Webcam Proctoring Telemetry: Incorporating WebRTC-based AI proctoring to detect multiple faces, unauthorized electronic devices, and user absence during critical exams.")
    bullet("Native Mobile Clients: Developing cross-platform Flutter/React Native mobile applications with offline lesson synchronization and secure exam sandbox environments.")

    # ==================== REFERENCES ====================
    h1("REFERENCES")
    refs = [
        "[1] PWD301 Development Team, 'PPTQ LMS – Intelligent Enterprise Learning & Assessment Platform,' GitHub Repository, 2026. [Online]. Available: https://github.com/concatapchay123/PWD301.",
        "[2] Pallets Projects, 'Flask Documentation (Version 3.x),' Pallets, 2026. [Online]. Available: https://flask.palletsprojects.com/.",
        "[3] Microsoft Corporation, 'SQL Server 2022 Technical Documentation & Transact-SQL Reference,' Microsoft Learn, 2026. [Online]. Available: https://learn.microsoft.com/sql/sql-server/.",
        "[4] Docker Inc., 'Docker Compose Specification and Orchestration Guide,' Docker Docs, 2026. [Online]. Available: https://docs.docker.com/compose/.",
        "[5] OWASP Foundation, 'OWASP Top 10: The Ten Most Critical Web Application Security Risks,' OWASP Foundation, 2025. [Online]. Available: https://owasp.org/www-project-top-ten/.",
        "[6] Google DeepMind, 'Google Gemini API Developer Guide & Python SDK Documentation,' Google Cloud, 2026. [Online]. Available: https://ai.google.dev/.",
        "[7] ClamAV Open Source Community, 'ClamAV Antivirus Daemon Specification and Client Protocol,' ClamAV Documentation, 2026. [Online]. Available: https://www.clamav.net/.",
        "[8] W3C, 'Web Content Accessibility Guidelines (WCAG) 2.2,' World Wide Web Consortium, 2024. [Online]. Available: https://www.w3.org/TR/WCAG22/."
    ]
    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.15
        r = p_ref.add_run(ref)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(51, 65, 85)

    # ==================== APPENDICES ====================
    h1("APPENDICES")
    
    # ---------------- Appendix A ----------------
    h2("Appendix A: Installation, Docker Orchestration, and Setup Guide")
    p_body(
        "This appendix provides full technical instructions for deploying PPTQ LMS from source code using Docker Compose [4]. These instructions guarantee environment reproducibility across Windows, Linux, and macOS."
    )
    h3("A.1. Environment Prerequisites")
    bullet("Operating System: 64-bit Windows 11, Ubuntu 22.04 LTS, or macOS Sonoma.")
    bullet("Docker Engine: Version 24.0.0 or higher with Docker Compose V2 [4].")
    bullet("System Memory: Minimum 8 GB RAM (16 GB recommended to support concurrent SQL Server and ClamAV containers).")
    bullet("Storage Capacity: Minimum 20 GB free disk space for container images, database files, and uploaded video assets.")

    h3("A.2. Environment Configuration (.env)")
    p_body("Create a production or development .env file at repository root with the parameters specified in .env.example:")
    callout(
        "FLASK_APP=wsgi.py\n"
        "FLASK_ENV=production\n"
        "SECRET_KEY=pwd301_production_super_secret_cryptographic_key_2026\n"
        "DATABASE_URL=mssql+pyodbc://sa:AdminStrongPassw0rd!@db:1433/pwd301_lms?driver=ODBC+Driver+18+for+SQL+Server&TrustServerCertificate=yes\n"
        "CLAMAV_HOST=clamav\n"
        "CLAMAV_PORT=3310\n"
        "GEMINI_API_KEY=AIzaSyA_key1,AIzaSyB_key2,AIzaSyC_key3\n"
        "MAX_CONTENT_LENGTH=1073741824",
        "Sample .env Environment Configuration"
    )

    h3("A.3. Docker Compose Orchestration Commands")
    p_body("Execute the following terminal commands to clone, configure, build, and initialize the system:")
    bullet("Step 1: Clone Repository: git clone https://github.com/concatapchay123/PWD301.git && cd PWD301")
    bullet("Step 2: Initialize Configuration: cp .env.example .env")
    bullet("Step 3: Build & Launch Containers: docker compose up -d --build")
    bullet("Step 4: Execute Database Migrations: docker compose exec web flask db upgrade")
    bullet("Step 5: Seed Base Data & Roles: docker compose exec web flask seed-baseline")
    bullet("Step 6: Seed Demo Fixtures (Optional): docker compose exec web flask seed-demo")

    h3("A.4. Service Status Verification")
    p_body("Verify container health using the docker compose ps command [4]. The operational status, healthcheck probes, and host port mappings across all four core containerized services are structured in Table A.1:")
    
    ps_headers = ["Service Name (Container)", "Docker Image Source", "Operational Health Status", "Host Port Mappings"]
    ps_rows = [
        ["pwd301_clamav (2f59f9a363d4)", "clamav/clamav:latest", "Up 9 hours (healthy)", "0.0.0.0:3310 -> 3310/tcp"],
        ["pwd301_db (f67cfa81fae7)", "mcr.microsoft.com/mssql/server:2022-latest", "Up 10 hours (healthy)", "0.0.0.0:1433 -> 1433/tcp"],
        ["pwd301_web", "pwd301-web:latest", "Up 45 minutes (healthy)", "0.0.0.0:5000 -> 5000/tcp"],
        ["pwd301_worker", "pwd301-web:latest", "Up 45 minutes (healthy)", "Internal Worker (No Host Port)"]
    ]
    tbl(ps_headers, ps_rows, [1.6, 2.1, 1.4, 1.4])
    p_body("Table A.1: Docker Compose Multi-Container Operational Healthcheck Status and Service Topology.", italic=True)

    h3("A.5. Accessing Application Instances")
    p_body("Once container health is confirmed, the web platform can be accessed at:")
    bullet("Local Development Instance: http://localhost:5000/ (REST API & Single-DOM SPA).")
    bullet("Live Institutional Production Instance: https://pptq.id.vn/ (Secured via TLS/HTTPS managed by Caddy).")

    # ---------------- Appendix B ----------------
    h2("Appendix B: Database Design, Schema Dictionary, and ERD")
    p_body(
        "The PPTQ database is deployed on Microsoft SQL Server 2022 [3]. This appendix outlines the structural domain distribution, provides an Entity-Relationship Diagram (ERD), and documents nine primary core entities directly corresponding to the SQLAlchemy models in src/pwd301/models/."
    )
    h3("B.1. Entity-Relationship Diagram (ERD Overview)")
    callout(
        "[users] (1) <---> (N) [courses] (via owner_instructor_id)\n"
        "[users] (1) <---> (N) [enrollments] (via student_user_id)\n"
        "[courses] (1) <---> (N) [learning_units] (via course_id)\n"
        "[learning_units] (1) <---> (N) [lessons] (via unit_id)\n"
        "[courses] (1) <---> (N) [assessments] (via course_id)\n"
        "[assessments] (1) <---> (N) [assessment_attempts] (via assessment_id)\n"
        "[users] (1) <---> (N) [assessment_attempts] (via student_user_id)\n"
        "[assessment_attempts] (1) <---> (N) [attempt_focus_events] (via attempt_id)\n"
        "[assessment_attempts] (1) <---> (1) [assessment_results] (via attempt_id)\n"
        "[users] (1) <---> (N) [audit_events] (via actor_user_id)",
        "Core Entity-Relationship Diagram (Structural Relationships)"
    )

    h3("B.2. Relational Domain Breakdown")
    domain_headers = ["Domain Name", "Primary Entity Models", "Key Responsibilities"]
    domain_data = [
        ["Identity & Access", "User, Role, UserRole, AuthSession", "Authentication, RBAC, session revocation"],
        ["Course & Curriculum", "Course, LearningUnit, Lesson, LessonResource", "Curriculum structure, lesson sequencing, attachments"],
        ["Assessment & Questions", "Question, QuestionChoice, AssessmentQuestion", "Question authoring, answer choice structuring"],
        ["Assessment Engine", "Assessment, AssessmentSection, GradingRule", "Exam blueprints, time limits, scoring constraints"],
        ["Attempt & Regrading", "AssessmentAttempt, AttemptFocusEvent, AssessmentResult", "Single active lease, focus telemetry, grading"],
        ["Files & Storage", "FileAsset, FileScan, QuarantineRecord", "ClamAV scan status, quarantine fail-closed"],
        ["AI RAG Knowledge", "KnowledgeChunk, ChunkEmbedding, AIChatSession", "Course grounding chunks, student Q&A history [6]"],
        ["Notification & Audit", "AuditEvent, Notification, BroadcastAnnouncement", "Append-only security audit log, system broadcasts"],
        ["System Operations", "DatabaseBackup, SystemHealthMetric, BackgroundJob", "Disaster recovery, hardware telemetry monitoring"],
    ]
    tbl(domain_headers, domain_data, [1.8, 2.2, 2.5])
    p_body("Table B.1: Relational Domain Distribution across Platform Entities.", italic=True)

    h3("B.3. Core Entity Data Dictionaries (9 Primary Tables)")
    p_body("Data dictionaries for the nine primary database entities verified directly against SQLAlchemy models in src/pwd301/models/:")

    # 1. users
    p_body("1. Table: users (Identity Entity - src/pwd301/models/identity.py)", bold_prefix="Entity 1: ")
    u_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    u_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal surrogate primary key"],
        ["public_id", "UNIQUEIDENTIFIER", "NOT NULL, UNIQUE", "Cryptographically random public GUID"],
        ["email", "NVARCHAR(320)", "NOT NULL, UNIQUE", "Academic login email identifier"],
        ["password_hash", "NVARCHAR(255)", "NOT NULL", "Werkzeug PBKDF2 password hash"],
        ["display_name", "NVARCHAR(150)", "NOT NULL", "Student or instructor display name"],
        ["status", "VARCHAR(24)", "NOT NULL, DEFAULT 'ACTIVE'", "Account status (ACTIVE, SUSPENDED)"],
        ["created_at", "DATETIME2(3)", "NOT NULL, UTC", "Account registration timestamp"],
    ]
    tbl(u_headers, u_data, [1.5, 1.5, 1.5, 2.0])

    # 2. roles
    p_body("2. Table: roles (Role Entity - src/pwd301/models/identity.py)", bold_prefix="Entity 2: ")
    r_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    r_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal role primary key"],
        ["code", "VARCHAR(32)", "NOT NULL, UNIQUE", "Role code (ROLE_STUDENT, ROLE_INSTRUCTOR, ROLE_ADMIN)"],
        ["name", "NVARCHAR(100)", "NOT NULL", "Human-readable role display title"],
    ]
    tbl(r_headers, r_data, [1.5, 1.5, 1.5, 2.0])

    # 3. courses
    p_body("3. Table: courses (Course Entity - src/pwd301/models/course.py)", bold_prefix="Entity 3: ")
    c_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    c_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal course identifier"],
        ["public_id", "UNIQUEIDENTIFIER", "NOT NULL, UNIQUE", "Public course GUID identifier"],
        ["course_code", "NVARCHAR(50)", "NOT NULL, UNIQUE", "Course code (e.g., PWD301, CS101)"],
        ["title", "NVARCHAR(200)", "NOT NULL", "Full course display title"],
        ["owner_instructor_id", "BIGINT", "FK -> users(id)", "Assigned managing instructor"],
        ["status", "VARCHAR(32)", "NOT NULL, DEFAULT 'DRAFT'", "Status: DRAFT, SUBMITTED, APPROVED, ARCHIVED"],
        ["price", "BIGINT", "NOT NULL, DEFAULT 0", "Tuition fee in VND (governed by admin audit)"],
    ]
    tbl(c_headers, c_data, [1.5, 1.5, 1.5, 2.0])

    # 4. learning_units
    p_body("4. Table: learning_units (Curriculum Chapter - src/pwd301/models/course.py)", bold_prefix="Entity 4: ")
    lu_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    lu_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal chapter identifier"],
        ["course_id", "BIGINT", "FK -> courses(id)", "Parent course reference"],
        ["title", "NVARCHAR(200)", "NOT NULL", "Chapter display title"],
        ["position", "INT", "NOT NULL", "Sequential sorting order within course"],
        ["deleted_at", "DATETIME2(3)", "NULLABLE", "Soft-deletion timestamp"],
    ]
    tbl(lu_headers, lu_data, [1.5, 1.5, 1.5, 2.0])

    # 5. lessons
    p_body("5. Table: lessons (Individual Lesson - src/pwd301/models/course.py)", bold_prefix="Entity 5: ")
    l_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    l_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal lesson identifier"],
        ["unit_id", "BIGINT", "FK -> learning_units(id)", "Parent chapter reference"],
        ["title", "NVARCHAR(200)", "NOT NULL", "Lesson display title"],
        ["position", "INT", "NOT NULL", "Sequential sorting order within chapter"],
        ["estimated_duration_minutes", "INT", "NULLABLE", "Estimated study time in minutes"],
        ["published_at", "DATETIME2(3)", "NULLABLE", "Publication timestamp"],
    ]
    tbl(l_headers, l_data, [1.5, 1.5, 1.5, 2.0])

    # 6. questions
    p_body("6. Table: questions (Assessment Question Entity - src/pwd301/models/question_bank.py)", bold_prefix="Entity 6: ")
    q_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    q_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal question primary key"],
        ["public_id", "UNIQUEIDENTIFIER", "NOT NULL, UNIQUE", "Public question GUID identifier"],
        ["difficulty", "VARCHAR(20)", "NOT NULL", "Difficulty level (EASY, MEDIUM, HARD)"],
        ["learning_objective", "NVARCHAR(500)", "NULLABLE", "Mapped curriculum learning objective"],
        ["deleted_at", "DATETIME2(3)", "NULLABLE", "Soft-deletion audit timestamp"],
    ]
    tbl(q_headers, q_data, [1.5, 1.5, 1.5, 2.0])

    # 7. assessments
    p_body("7. Table: assessments (Exam Blueprint - src/pwd301/models/assessment.py)", bold_prefix="Entity 7: ")
    as_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    as_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal assessment primary key"],
        ["title", "NVARCHAR(200)", "NOT NULL", "Assessment display title"],
        ["assessment_type", "VARCHAR(20)", "NOT NULL", "Type: QUIZ, MIDTERM, FINAL"],
        ["time_limit_minutes", "INT", "NULLABLE", "Allotted examination duration"],
        ["attempt_limit", "INT", "NULLABLE", "Maximum permitted attempts per student"],
        ["passing_percent", "NUMERIC(5,2)", "NOT NULL, DEFAULT 70.0", "Passing score threshold percentage"],
        ["monitoring_enabled", "BIT", "NOT NULL, DEFAULT 1", "Enables blur & focus telemetry tracking"],
    ]
    tbl(as_headers, as_data, [1.5, 1.5, 1.5, 2.0])

    # 8. assessment_attempts
    p_body("8. Table: assessment_attempts (Exam Attempt - src/pwd301/models/attempt_regrade.py)", bold_prefix="Entity 8: ")
    aa_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    aa_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal attempt surrogate key"],
        ["public_id", "UNIQUEIDENTIFIER", "NOT NULL, UNIQUE", "Public attempt GUID identifier"],
        ["assessment_id", "BIGINT", "FK -> assessments(id)", "Associated examination blueprint"],
        ["student_user_id", "BIGINT", "FK -> users(id)", "Candidate student taking the exam"],
        ["attempt_number", "INT", "NOT NULL", "Monotonically increasing attempt count"],
        ["status", "VARCHAR(28)", "NOT NULL", "Status: IN_PROGRESS, SUBMITTED, EVALUATED"],
        ["lease_token_hash", "BINARY(32)", "NULLABLE", "SHA-256 hash of active editing lease token"],
        ["lease_epoch", "INT", "NOT NULL, DEFAULT 1", "Sequential epoch token preventing stale writes"],
        ["lease_expires_at", "DATETIME2(3)", "NULLABLE", "Lease expiry deadline (UTC)"],
    ]
    tbl(aa_headers, aa_data, [1.5, 1.5, 1.5, 2.0])

    # 9. audit_events
    p_body("9. Table: audit_events (System Audit Trail - src/pwd301/models/notification_audit.py)", bold_prefix="Entity 9: ")
    ae_headers = ["Column Name", "Data Type", "Constraints", "Description"]
    ae_data = [
        ["id", "BIGINT", "PK, IDENTITY(1,1)", "Internal audit event identifier"],
        ["actor_user_id", "BIGINT", "FK -> users(id)", "User triggering the audited action"],
        ["action", "VARCHAR(80)", "NOT NULL", "Audited action (e.g., COURSE_APPROVED, USER_SUSPENDED)"],
        ["target_type", "VARCHAR(40)", "NOT NULL", "Target entity type (e.g., COURSE, USER, ATTEMPT)"],
        ["target_id", "BIGINT", "NULLABLE", "Surrogate ID of the target entity"],
        ["reason", "NVARCHAR(1000)", "NULLABLE", "Administrative reason for action"],
        ["created_at", "DATETIME2(3)", "NOT NULL, UTC", "Append-only immutable event timestamp"],
    ]
    tbl(ae_headers, ae_data, [1.5, 1.5, 1.5, 2.0])

    # ---------------- Appendix C ----------------
    h2("Appendix C: Test Cases, Concurrency Stress, and Verification Matrix")
    p_body(
        "This appendix documents empirical test evidence, boundary stress tests, and concurrency race condition evaluations conducted across the platform."
    )
    h3("C.1. Automated Test Suite Metrics")
    bullet("Test Suite Structure: Structured across tests/unit/, tests/integration/, and tests/e2e/ executed via pytest.")
    bullet("Core Domain Coverage: Validates identity, curriculum lifecycle, question authoring, assessment leasing, and grading.")
    bullet("API Contract Conformance: Verifies standardized JSON envelopes across authentication, student, instructor, and admin routes.")
    bullet("Browser Scenarios: Automated Playwright user workflows verifying catalog browsing, video playback, and exam submission.")

    h3("C.2. Critical Invariant Verification Matrix")
    inv_headers = ["Security & Logic Invariant", "Evaluation Scenario", "Observed System Behavior", "Audit Verdict"]
    inv_data = [
        ["Single Active Lease (Alg 07)", "Student opens active exam in a second browser tab", "Tab 2 blocked with HTTP 409 (LEASE_CONFLICT); Tab 1 continues", "Pass (Blocked)"],
        ["50MB Non-Archive Attachment", "Instructor attempts to upload sample.zip (2.4MB)", "Rejected with HTTP 400 & institutional Vietnamese notice", "Pass (Enforced)"],
        ["Fail-Closed ClamAV Scan", "Upload EICAR standard test virus signature file", "Quarantined immediately; status set to QUARANTINED", "Pass (Fail-Closed)"],
        ["5-Second Pacing Unlock", "Student scrubs to (Duration - 4s) in video", "Player permits seek; is_video_complete credited", "Pass (Unlocked)"],
        ["Anti-Tamper Lockdown", "Candidate opens DevTools during active exam", "Session revoked, blackout security shield triggered", "Pass (Locked Down)"],
        ["Gradebook PDF Evaluation", "Student scores 5.0/10.0 (below 7.0 standard)", "PDF generates with bold red 'KHÔNG ĐẠT' header", "Pass (Compliant)"],
    ]
    tbl(inv_headers, inv_data, [1.6, 1.8, 2.0, 1.1])
    p_body("Table C.1: Verification Matrix for Critical Architectural Invariants.", italic=True)

    # ---------------- Appendix D ----------------
    h2("Appendix D: User Manual and Interface Evidence")
    p_body(
        "This appendix provides operational walkthroughs for Students, Instructors, and Administrators, supported by photographic evidence captured from the live deployment at https://pptq.id.vn/."
    )

    h3("D.1. User Authentication and Login Experience")
    p_body(
        "Users navigate to the login portal (https://pptq.id.vn/#/auth) and provide their registered academic email and password. Upon successful verification, the server establishes a secure session cookie and redirects the client to their role-specific portal."
    )
    fig(r'E:/PWD301/scripts/extracted_images/image3.png', "Figure D.1: Login Interface of the PPTQ System (https://pptq.id.vn/#/auth).")

    h3("D.2. Student Course Discovery and Learning Console")
    p_body(
        "Students access their enrolled courses or browse the institutional catalog. The learning console provides a structured environment featuring an outline navigation tree, lesson markdown reader, and protected 16:9 video player."
    )
    fig(r'E:/PWD301/.agents/teamwork/evidence_r6/r09_course_intro.png', "Figure D.2: Student Course Catalog and Introduction Interface.")
    # Use r15_local_video_play.png to show actual working video playback
    fig(r'E:/PWD301/.agents/teamwork/evidence_r6/r15_local_video_play.png', "Figure D.3: Protected Video Learning Player and Curriculum Outline Interface.")

    h3("D.3. Instructor Studio and Course Authoring")
    p_body(
        "Instructors utilize the Studio dashboard to manage their teaching portfolio, structure chapters, attach learning documents, and author exams."
    )
    fig(r'E:/PWD301/.agents/teamwork/evidence_r6/r02_instructor_dashboard.png', "Figure D.4: Instructor Studio Management Dashboard.")
    fig(r'E:/PWD301/scripts/extracted_images/image4.png', "Figure D.5: Course Creation Modal Interface with Instructor Advisory Notice.")

    h3("D.4. Administrator Governance and Review Portal")
    p_body(
        "System administrators supervise institutional operations, review staged course changesets, audit price modifications, and monitor server hardware telemetry."
    )
    fig(r'E:/PWD301/.superpowers/task-090/browser-admin-preview.png', "Figure D.6: Administrative Course Review and Inspection Interface.")
    # Use admin-review-history-desktop.png to show actual approval history & audit events
    fig(r'E:/PWD301/.superpowers/task095\admin-review-history-desktop.png', "Figure D.7: Administrative Approval Workflow and Course Audit History Interface.")

    h3("D.5. Server-Authoritative Online Examination and Results")
    p_body(
        "During online examinations, candidates operate within an enforced fullscreen console. The anti-cheat engine tracks window blur events and remaining time in real-time."
    )
    fig(r'E:/PWD301/scripts/extracted_images/image5.png', "Figure D.8: Online Examination Interface featuring Focus Monitoring and UTC Timer.")
    p_body(
        "Upon submission, answers are graded instantaneously. Candidates review their question breakdown, examine detailed academic explanations, request appeals, or export official PDF transcripts."
    )
    fig(r'E:/PWD301/scripts/extracted_images/image6.png', "Figure D.9: Detailed Examination Results Interface with Academic Feedback and PDF Export.")

    # Save Word Document
    doc.save(output_docx)
    print(f"Generated DOCX successfully: {output_docx}")

    # Generate Markdown Document
    print(f"Generating Markdown version: {output_md}...")
    generate_markdown_version(output_md)
    print("Done generating all reports!")

def generate_markdown_version(output_md):
    content = """# ONLINE COURSE MANAGEMENT PLATFORM (PPTQ LMS)
## Final Report — PWD301 (Python Web Development)
**FPT University — Fall 2026 — Class IA2007 — Group 1**  
**Academic Supervisor:** Lê Tiến Dũng  
**Development Team:**
- SE203140 - Đặng Lý Quân
- SE201470 - Lại Vĩnh Phú
- SE204474 - Trần Đặng Hữu Thắng
- SE203815 - Phạm Nguyễn Hoàng Phúc

**Production Deployment:** https://pptq.id.vn/  
**Source Code Repository:** https://github.com/concatapchay123/PWD301

---

## ABSTRACT
With the rapid advancement of modern information technology, web-based learning and assessment platforms have become critical infrastructure for academic institutions and modern enterprises [1]. However, existing commercial and open-source learning management systems often present significant architectural shortcomings, including tightly coupled server-side templating engines, unverified client-side examination timers vulnerable to local tampering, and inadequate fail-open file handling mechanisms that expose systems to malware. To address these operational challenges, our engineering team designed and developed **PPTQ LMS**—a 100% custom-built, enterprise-grade Online Course Management and Assessment Platform engineered from the ground up for the PWD301 (Python Web Development) course at FPT University.

The PPTQ platform is architected around a high-performance **Modular Monolith with a Pure Headless REST API** powered by Python 3.11+ and Flask 3.x [2], decoupling server-side business logic from the user presentation layer. All domain data is managed through an enterprise-grade **Microsoft SQL Server 2022** database [3] structured around normalized relational entities enforcing ACID transactions, strict foreign key constraints, and optimistic concurrency controls. The client presentation layer utilizes a lightweight **Single-DOM Single-Page Application (SPA)** built in modern Vanilla JavaScript (ES2022+) and styled with Tailwind CSS alongside KaTeX mathematical typography, eliminating heavy frontend framework overhead while ensuring responsive client transitions.

Key architectural mechanisms implemented within PPTQ include:
1. **Server-Authoritative Online Examination Engine**: Enforces a Single Active Editing Lease algorithm (Algorithm 07) that strictly blocks secondary tab conflicts, coupled with client-side focus monitoring (blur and tab-switch telemetry).
2. **Anti-Forward Seek Video Pacing Engine**: Automated 5-second completion unlock and real-time lesson progression without page reload.
3. **Fail-Closed Antivirus Security Pipeline**: ClamAV daemon antivirus scanning coupled with a strict 50 MB Non-Archive Attachment Policy [7].
4. **Institutional AI Assistant ('Bạch Tuột')**: Powered by Google Gemini [6] with multi-key pool rotation, domain-grounded retrieval, and interactive question-level review explanations.

The platform is containerized via Docker Compose [4] and verified live on production at https://pptq.id.vn/.

*Keywords: Online Course Management System, Learning Management System, Flask, Python, Microsoft SQL Server, Docker Compose, Headless REST API, Server-Authoritative Assessment, AI Assistant, Single-DOM SPA.*

---

## I. INTRODUCTION

### 1.1. Background and Motivation
The digital transformation of global education has demonstrated that online learning systems are foundational environments for knowledge dissemination, competency assessment, and academic governance. Despite the proliferation of legacy platforms, institutional educators consistently face deep architectural and operational challenges. Traditional monoliths often tightly couple UI rendering templates with core business services, leading to maintainability bottlenecks, bloated page payloads, and severe security risks during concurrent exam sessions [1], [5].

Furthermore, academic integrity remains a vulnerable touchpoint in online education. Many existing solutions rely on client-side JavaScript timers for exam countdowns, which can be manipulated using browser developer tools, or permit unrestricted multi-tabbing that enables unmonitored collusion. Similarly, unverified file uploads represent a severe attack vector, where malicious archives or ransomware can be injected under the guise of study materials [5], [7].

Recognizing these critical systemic gaps, our team undertook the development of PPTQ: an Online Course Management and Assessment Platform. Developed as the final capstone for PWD301 – Python Web Development at FPT University, PPTQ was engineered 100% in-house from scratch. The project provides an empirical demonstration of how modern Python software engineering principles, rigorous database normalization, fail-closed security protocols, and human-centered interface design coalesce into a robust academic platform.

### 1.2. Project Objectives
#### 1.2.1. General Objective
The overarching objective of the project is to architect, develop, verify, and deploy a comprehensive, production-grade web application supporting online course management, curriculum organization, protected video learning, secure online examination, and AI-powered academic assistance across three distinct user roles: Students, Instructors, and System Administrators.

#### 1.2.2. Specific Objectives
- Engineer a modular, headless backend architecture using Python 3.11 and Flask 3.x, delivering standardized JSON REST API envelopes across all operational domains [2].
- Implement a high-performance Single-DOM Single-Page Application (SPA) delivering seamless navigation, instantaneous state updates, and WCAG AA accessibility compliance [8].
- Design and normalize an enterprise-grade database schema on Microsoft SQL Server 2022 comprising relational tables with strict foreign keys, BIGINT identities, public GUIDs, and UTC timestamps [3].
- Develop a server-authoritative online examination engine featuring the Single Active Editing Lease algorithm (Algorithm 07), real-time focus/tab-switch tracking, idempotent submissions, and automated grading.
- Implement an anti-forward seek video pacing engine with HLS streaming security and a 5-second completion unlock mechanism.
- Enforce a fail-closed file security architecture incorporating real-time ClamAV daemon antivirus scanning and a strict 50 MB Non-Archive Attachment Policy [7].
- Integrate 'Bạch Tuột'—an institutional AI Assistant powered by Google Gemini with multi-key pool rotation, retrieval-augmented grounding, and dedicated question-level review explanations [6].
- Execute comprehensive automated testing across unit, integration, and security test suites, orchestrating all services via multi-container Docker Compose [4].

### 1.3. Project Scope
The project scope encompasses the complete software development lifecycle, from domain modeling and API specification to containerized deployment. The functional boundaries are explicitly delineated across three core operational roles:
- **Student Domain**: Course discovery, permanent enrollment, structured learning view with protected video streaming, document downloads, interactive mini-quizzes, anti-cheat online examination, and academic gradebook transcript generation.
- **Instructor Domain**: Studio course creation, curriculum outline structuring (Chapters and Lessons), 50 MB document attachment uploads, exam authoring across standardized question formats (Single Choice, Multiple Choice, True/False), assessment scheduling, and unified staging changeset approvals.
- **Administrator Domain**: System-wide user account lifecycle governance (activation, suspension, and role assignment), course publication review and rejection workflows, fee governance with audit trail logging, and live hardware telemetry monitoring via `psutil`.

> **System Invariant**: PPTQ LMS is an independent, 100% custom-built platform engineered by Group 1. It contains zero legacy code from third-party commercial platforms and operates entirely on dedicated, containerized services.

### 1.4. Target Users
- **Students (Learners)**: Require an intuitive, distraction-free learning environment that tracks progress accurately, provides responsive video playback, protects their examination submissions against connectivity drops, and offers immediate academic support.
- **Instructors (Educators)**: Require a cohesive management studio to build rich modular curricula, author and manage examinations with customizable constraints (time limits, attempt limits, pass rates), and analyze student grading outcomes transparently.
- **System Administrators (Governance)**: Require granular role-based access control, real-time server health metrics, tamper-evident audit logs, and centralized course review workflows to guarantee institutional compliance.

### 1.5. Development Methodology
- **Phase 1 - Domain Modeling & Specification**: Defined canonical business rules, data dictionaries, state machines, and API contracts before authoring application code.
- **Phase 2 - Test-Driven Implementation (TDD Iron Law)**: Authored automated failing tests (unit, integration, and security boundaries) prior to implementing backend route handlers or services.
- **Phase 3 - Modular Monolith & Headless Separation**: Built discrete Flask blueprints with centralized error handling and request validation, strictly isolating business logic from presentation concerns.
- **Phase 4 - Verification Before Completion**: Executed automated test pipelines, real-browser Playwright suites, and security checks before marking features as complete.
- **Phase 5 - Multi-Container Orchestration**: Packaged the application, database, and security services into Docker Compose, validating production parity on https://pptq.id.vn/ [4].

---

## II. SYSTEM REQUIREMENTS ANALYSIS

### 2.1. Problem Overview
Modern educational environments require centralized platforms capable of managing complex educational hierarchies while maintaining strict data integrity. Key challenges that PPTQ was engineered to resolve include:
- **Data Fragmentation**: Consolidating disparate modules for identity, video streaming, document repository, and online examination assessments into a unified relational core.
- **Exam Integrity Vulnerabilities**: Eliminating client-side timer manipulation, concurrent multi-tab cheating, and unsaved progress loss during network interruptions.
- **Security & Malware Exposure**: Protecting institutional servers against malicious file uploads and zip-bomb attacks via automated fail-closed scanning [5], [7].
- **AI Reliability & Cost Control**: Preventing AI hallucination and service outages through retrieval grounding and multi-key API pool load balancing [6].

### 2.2. Functional Requirements
- **FR-01: User Authentication & Identity Lifecycle**: Secure registration, login via academic email, Werkzeug PBKDF2 password verification, Flask session cookie management with HttpOnly and SameSite=Lax flags, JWT issuance for external API consumers, and instantaneous session revocation upon suspension.
- **FR-02: Role-Based Access Control (RBAC)**: Strict server-side authorization enforcing object-level ownership checks across Student, Instructor, and Administrator roles.
- **FR-03: Course Lifecycle & Staging Drafts**: Unified staging draft aggregate allowing instructors to stage content modifications continuously and submit them atomically for administrative review.
- **FR-04: Hierarchical Curriculum Management**: Structuring learning content into multi-tiered Course -> LearningUnit (Chapter) -> Lesson hierarchies with drag-and-drop sequencing.
- **FR-05: Protected Video Learning & Pacing Engine**: Encrypted HLS video streaming with session-bound 12-hour DRM tokens, standard 16:9 responsive viewport, anti-forward seek restriction up to (Duration - 5s), 5-second completion unlock, and real-time next-lesson gate.
- **FR-06: Server-Authoritative Online Examination**: Fullscreen exam console, focus/tab-switch blur telemetry counter, Single Active Editing Lease algorithm (Algorithm 07), and server-managed UTC countdown timers.
- **FR-07: Grading, Regrading & Academic Transcripts**: Idempotent submission processing, automated question scoring, audit-trailed manual regrading, and official PDF gradebook generation with 'KHÔNG ĐẠT' / 'ĐẠT' evaluation labels.
- **FR-08: System Governance & Authentic Telemetry**: Real-time hardware telemetry (CPU, RAM, Disk) queried via `psutil`, course fee governance, and append-only audit event logging.
- **FR-09: Intelligent Octopus AI Assistant**: Domain-grounded question answering via Google Gemini, strict course-scope RAG retrieval, multi-key rotating pool, and interactive question-level review explanations [6].

### 2.3. Non-Functional Requirements
- **Security**: Defense-in-depth architecture adhering to OWASP Top 10 standards [5], strict CSRF token validation, input sanitization, and fail-closed ClamAV file scanning [7].
- **Data Integrity**: ACID transactions on Microsoft SQL Server 2022 [3], strict foreign keys, BIGINT internal keys, GUID public exposures, and rowversion concurrency tokens.
- **Usability**: Warm Editorial design system adhering to Jakob's, Hick's, and Miller's psychological laws, single-column forms, skeleton loading states, and WCAG AA contrast ratio (>= 4.5:1) [8].
- **Maintainability**: Modular Monolith architecture with clean blueprint separation, modular serialization schemas, and type-annotated domain logic.
- **Deployability**: Fully reproducible, multi-container Docker Compose deployment [4] guaranteeing zero configuration drift between development and production environments.
- **Performance & Reliability**: Sub-100ms API response latency for cached endpoints, automated container healthcheck probes, and high availability targets.

### 2.4. Role-Based Access Control Matrix
| Functional Capability | Student Role | Instructor Role | Administrator Role |
| :--- | :--- | :--- | :--- |
| User Registration & Profile Update | Self Only | Self Only | Full Management |
| Course Enrollment (Lifetime) | Self Enrolled | Assigned Courses | System-wide Override |
| Curriculum & Lesson Authoring | Read (Enrolled) | Create / Edit Owned | Review / Approve / Delete |
| Document Attachment Upload (<50MB) | Forbidden | Allowed (Owned) | Full Management |
| Video Pacing & Lesson Completion | Personal Tracker | Preview Mode | Audit / Inspection |
| Exam & Question Authoring | Forbidden | Create / Edit Owned | System-wide Oversight |
| Exam Creation & Scheduling | Forbidden | Create / Schedule Owned | Audit & Emergency Override |
| Exam Attempt & Submission | Allowed (Enrolled) | Preview Mode Only | Forbidden (Audit Only) |
| Score Regrading & Appeal | View & Request Appeal | Execute Regrade (Owned) | Final Arbitration |
| System Telemetry & Server Logs | Forbidden | Forbidden | Full Hardware Telemetry |

---

## III. SYSTEM DESIGN

### 3.1. Overall Architecture
PPTQ LMS is designed as a four-tier Modular Monolith combining monolithic deployment simplicity with strict headless API boundaries:
1. **Tier 1 - Presentation Layer**: Single-DOM SPA written in modern Vanilla JavaScript (ES2022+) styled with Tailwind CSS (via CDN), KaTeX math rendering, and Google Fonts. Communicates exclusively via asynchronous JSON fetch requests.
2. **Tier 2 - Security Gateway**: Host-level Caddy reverse proxy providing automatic HTTPS/TLS termination, rate-limiting, and security headers. Session state maintained via signed, HttpOnly cookies with CSRF tokens.
3. **Tier 3 - Application Core**: Flask Modular Monolith with blueprints [2] for Auth, Student Experience, Instructor Studio, Admin Governance, Assessment Engine, and AI Assistant.
4. **Tier 4 - Persistence & Containers**: Docker-managed multi-container architecture [4] comprising Microsoft SQL Server 2022 [3] (`db`), ClamAV daemon (`clamav`), background worker (`worker`), and Flask web API (`web`), with host hardware telemetry queried via `psutil`.

### 3.2. Database Design (Relational Architecture)
The database schema on Microsoft SQL Server 2022 [3] is organized into nine functional domains:
- **Identity & Access**: `users`, `roles`, `user_roles`, `auth_sessions`, `login_history`, `password_resets`, `permissions`, `role_permissions`.
- **Course & Curriculum**: `courses`, `learning_units`, `lessons`, `course_prerequisites`, `course_prices`, `categories`, `course_instructors`, `enrollments`, `lesson_resources`, `lesson_completions`, `video_trackers`, `course_versions`.
- **Assessment & Questions**: Question items, answer choices, revisions, matching pairs, and assessment section blueprints supporting single-choice, multiple-choice, and true/false evaluation.
- **Assessment Engine**: `assessments`, `assessment_sections`, `assessment_questions`, `assessment_policies`, `assessment_schedules`, `assessment_access_codes`, `grading_rules`, `regrade_policies`, `rubrics`, `rubric_criteria`.
- **Attempt & Regrading**: `assessment_attempts`, `attempt_answers`, `attempt_focus_events`, `assessment_results`, `assessment_result_history`, `regrade_items`, `score_histories`, `assessment_appeals`, `item_responses`, `attempt_questions`, `certificates`.
- **Files & Storage**: `file_assets`, `file_scans`, `quarantine_records`, `upload_sessions`, `storage_nodes`, `lesson_file_mappings`, `temp_uploads`.
- **AI RAG Knowledge**: `knowledge_chunks`, `chunk_embeddings`, `ai_chat_sessions`, `ai_messages`, `ai_feedback`, `grounding_references` [6].
- **Notification & Audit**: `audit_events`, `notifications`, `user_notifications`, `broadcast_announcements`, `audit_snapshots`, `system_alerts`, `email_outbox`, `webhook_subscriptions`.
- **Operations & Telemetry**: `database_backups`, `restore_history`, `system_health_metrics`, `telemetry_snapshots`, `background_jobs`, `maintenance_locks`.

### 3.3. User Interface Design
The user interface is governed by the Warm Editorial aesthetic system:
- **Jakob's Law**: Standardized topbar, sidebar, and breadcrumb navigation.
- **Hick's Law**: Clean single-column forms with explicit call-to-actions.
- **Miller's Law (7±2)**: Information chunking across course outlines and question navigation matrices.
- **Clean History Stack**: Dynamic routing using `history.replaceState` eliminates back-button navigation traps on form creation views.

---

## IV. TECHNOLOGIES AND DEPLOYMENT

### 4.1. Technology Stack Summary
- **Backend Language & Framework**: Python 3.11+, Flask 3.x [2], SQLAlchemy 2.0 ORM with `pyodbc` [3].
- **Database Engine**: Microsoft SQL Server 2022 (Docker Linux container `pwd301_db`, port 1433) [3].
- **Frontend Stack**: Vanilla JavaScript (ES2022+), Tailwind CSS CDN, KaTeX math typesetting, HTML5 APIs.
- **Security & Antivirus**: ClamAV Antivirus Daemon (port 3310) [7], Werkzeug PBKDF2 password hashing.
- **AI Engine**: Google Gemini API with multi-key rotation pool and RAG retrieval [6].
- **Containerization**: Docker & Docker Compose V2 [4] (4 services: `db`, `clamav`, `web`, `worker`).
- **Production Server**: Ubuntu 22.04 LTS host with Caddy reverse proxy at `https://pptq.id.vn/`.

---

## V. FEATURE IMPLEMENTATION

### 5.1. Authentication and Identity Governance
- **Session Authentication**: Secure, HttpOnly, SameSite=Lax cookies with CSRF defense for Web SPA.
- **JWT Authentication**: Short-lived Bearer tokens for external REST API endpoints.
- **Instant Revocation**: Account suspension or credential update immediately increments `auth_version`, invalidating active session cookies and cached tokens.

### 5.2. Course and Curriculum Management
- **50 MB Non-Archive Policy**: Strict rejection of `.zip`, `.rar`, `.7z` with standardized Vietnamese alert message.
- **Fail-Closed Antivirus**: Files pending scan or quarantined are blocked from student download [7].
- **Protected Video Streaming**: HLS streaming with 12-hour session DRM tokens, 16:9 responsive viewport, and anti-forward seek up to (Duration - 5s).
- **5-Second Seek Unlock**: Free scrubbing enabled in final 5 seconds; automatic lesson completion credit.
- **Real-Time Next Lesson Gate**: 'Tiếp theo' button dynamically lights up green (`bg-emerald-600`) upon completion without requiring page reload.

### 5.3. Server-Authoritative Online Examination Engine
- **Single Active Editing Lease (Algorithm 07)**: Only 1 active tab permitted per attempt via `lease_token_hash` and `lease_epoch`. Opening a second tab raises an `AttemptLeaseConflictError` (HTTP 409 Conflict); takeover is only allowed after expiry or explicit release.
- **Anti-Cheat Focus Tracking**: Fullscreen requested on start; window blur and fullscreen exit logged to `attempt_focus_events` table with telemetry timer.
- **Server UTC Countdown**: Eliminates local device clock manipulation vulnerabilities.
- **Idempotent Submission**: Retries and duplicate network requests never corrupt recorded exam scores.
- **Anti-Tamper Lockdown**: Opening DevTools / F12 triggers immediate session revocation and blackout security alert.
- **Official PDF Transcript**: Generates official PDF gradebook with bold red 'KHÔNG ĐẠT' for non-passing attempts.

### 5.4. Intelligent AI Assistant (Bạch Tuột)
- **Official Identity**: "Bạch Tuột, trợ lí AI của hệ thống PPTQ".
- **Interactive Assistant Touchpoints**:
  1. *Floating AI Tutor*: Greets students upon opening with 'Xin chào bạn! Mình là Bạch Tuột, trợ lí AI của hệ thống PPTQ. Mình có thể giúp được gì cho bạn?' providing contextual study guidance.
  2. *Instant Post-Exam Review ('Hỏi Bạch tuộc câu này')*: Embedded directly alongside assessment results (Figure D.9), enabling candidates to receive immediate, question-specific pedagogical explanations for why choices were correct or incorrect.
  3. *Dynamic Scope Guardrails*: Gracefully filters off-topic or adversarial prompts to keep interactions focused on academic coursework.
- **RAG Grounding**: Answers grounded strictly in active course syllabus; designed to avoid hallucination [6].
- **Defensive Prompt Filtering**: Sanitizes user queries to reject adversarial attempts to extract internal database records [5].

---

## VI. TESTING AND EVALUATION

### 6.1. Core Functional Test Execution Results
| Test ID | Test Scenario Description | Expected Result | Actual Result | Latency | Status | Notes & Observations |
| :--- | :--- | :--- | :--- | :--- | :--- | :--- |
| **TC-01** | User Login with Academic Email | Session cookie issued, redirect to role portal | Redirected, valid session cookie | 42ms | **Pass (100%)** | Werkzeug PBKDF2 verification verified |
| **TC-02** | Login with Invalid Password | HTTP 401, localized Vietnamese error envelope | HTTP 401 returned, error rendered | 28ms | **Pass (100%)** | Invalid attempt counter incremented (Brute-force defense) |
| **TC-03** | Student Accesses Admin Endpoint | HTTP 403 Forbidden, request blocked | HTTP 403 returned, audit logged | 15ms | **Pass (100%)** | RBAC guardrails functioning |
| **TC-04** | Course Outline & Catalog Navigation | Complete hierarchical curriculum returned | Hierarchy matches 3NF schema | 35ms | **Pass (100%)** | Hierarchy matches 3NF schema |
| **TC-05** | Online Exam Submission & Grading | Server scores answers, attempt locked | Idempotent submit, score recorded | 68ms | **Pass (100%)** | Lease successfully released |
| **TC-06** | PDF Gradebook Transcript Export | PDF downloaded, 'KHÔNG ĐẠT' rendered in red | PDF downloaded, valid watermark | 120ms | **Pass (100%)** | ReportLab generator active |
| **TC-07** | Octopus AI Knowledge Retrieval | Accurate lesson explanation, guardrails enforced | Accurate response, injection blocked | 850ms | **Pass (100%)** | Gemini pool active, Injection defense verified |
| **TC-08** | Docker Multi-Container Healthcheck | All core containers report Up (healthy) | All services healthy and running | N/A | **Pass (100%)** | Zero container crashes observed |

### 6.2. Test Coverage and Verification Summary
The test suite spans unit tests (`tests/unit/`), integration tests (`tests/integration/`), and end-to-end browser workflows (`tests/e2e/`), systematically validating business rules, database transactions, and security boundaries.

---

## VII. CONCLUSION AND FUTURE WORK
PPTQ LMS demonstrates that an in-house, Python/Flask-based architecture backed by Microsoft SQL Server 2022 and Docker Compose can deliver reliable performance, examination integrity, and modern educational UX. Operating live at https://pptq.id.vn/, the platform fulfills all technical requirements of the PWD301 curriculum at FPT University.

Future expansions will focus on automated multi-bitrate HLS video transcoding queues, WebRTC-based AI webcam proctoring, and cross-platform native mobile applications.

---

## REFERENCES
- [1] PWD301 Development Team, 'PPTQ LMS – Intelligent Enterprise Learning & Assessment Platform,' GitHub Repository, 2026. [Online]. Available: https://github.com/concatapchay123/PWD301.
- [2] Pallets Projects, 'Flask Documentation (Version 3.x),' Pallets, 2026. [Online]. Available: https://flask.palletsprojects.com/.
- [3] Microsoft Corporation, 'SQL Server 2022 Technical Documentation & Transact-SQL Reference,' Microsoft Learn, 2026. [Online]. Available: https://learn.microsoft.com/sql/sql-server/.
- [4] Docker Inc., 'Docker Compose Specification and Orchestration Guide,' Docker Docs, 2026. [Online]. Available: https://docs.docker.com/compose/.
- [5] OWASP Foundation, 'OWASP Top 10: The Ten Most Critical Web Application Security Risks,' OWASP Foundation, 2025. [Online]. Available: https://owasp.org/www-project-top-ten//.
- [6] Google DeepMind, 'Google Gemini API Developer Guide & Python SDK Documentation,' Google Cloud, 2026. [Online]. Available: https://ai.google.dev/.
- [7] ClamAV Open Source Community, 'ClamAV Antivirus Daemon Specification and Client Protocol,' ClamAV Documentation, 2026. [Online]. Available: https://www.clamav.net/.
- [8] W3C, 'Web Content Accessibility Guidelines (WCAG) 2.2,' World Wide Web Consortium, 2024. [Online]. Available: https://www.w3.org/TR/WCAG22/.

---

## APPENDICES

### Appendix A: Installation, Docker Orchestration, and Setup Guide

#### A.4. Service Status Verification
Verify container health using the `docker compose ps` command [4]. The operational status, healthcheck probes, and host port mappings across all four core containerized services are structured below:

| Service Name (Container) | Docker Image Source | Operational Health Status | Host Port Mappings |
| :--- | :--- | :--- | :--- |
| **pwd301_clamav** (`2f59f9a363d4`) | `clamav/clamav:latest` | **Up 9 hours (healthy)** | `0.0.0.0:3310 -> 3310/tcp` |
| **pwd301_db** (`f67cfa81fae7`) | `mcr.microsoft.com/mssql/server:2022-latest` | **Up 10 hours (healthy)** | `0.0.0.0:1433 -> 1433/tcp` |
| **pwd301_web** | `pwd301-web:latest` | **Up 45 minutes (healthy)** | `0.0.0.0:5000 -> 5000/tcp` |
| **pwd301_worker** | `pwd301-web:latest` | **Up 45 minutes (healthy)** | Internal Worker (No Host Port) |

*Table A.1: Docker Compose Multi-Container Operational Healthcheck Status and Service Topology.*

- **Appendix B**: Database Design, ERD Overview, and 9 Core Entity Data Dictionaries (`users`, `roles`, `courses`, `learning_units`, `lessons`, `questions`, `assessments`, `assessment_attempts`, `audit_events`).
- **Appendix C**: Test Cases, Concurrency Stress, and Invariant Verification Matrix.
- **Appendix D**: Illustrated User Manual with Photographic Production Evidence (Figures D.1 to D.9).
"""
    with open(output_md, 'w', encoding='utf-8') as f:
        f.write(content)

if __name__ == '__main__':
    build_report()
